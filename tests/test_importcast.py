"""import-cast (P09 п.1): таблица персонажей .xlsx → content/characters.yaml.

Таблицы генерируются openpyxl по описанию формата (версия от 11.10.2026); контент — копия
content/ в tmp_path, настоящий content/ не меняется.
"""

from __future__ import annotations

import difflib
import re
import shutil
from pathlib import Path

import openpyxl
import pytest
import yaml

from engine.__main__ import main
from engine.content import CONTENT_DIR, ContentError, load_content
from engine.importcast import apply_edits, format_report, import_cast
from engine.names import patronymic_gender

CAST_HEADERS = ["№", "  роль (рабочее  имя) ", "Комментарий", "ФИО в документах", "ВОЗРАСТ в игре"]
BG_HEADERS = ["Кто", "Где упоминается", "Имя (заполни)", "Можно родственника?"]
EXAMPLE_CAST = ["0", "Пример роли", "пример", "Иванов Иван Иванович", 40]


@pytest.fixture()
def content(tmp_path: Path) -> Path:
    """Копия content/, где заглушки фоновых имён уже заполнены первым импортом:
    таблица из identity_rows() для неё — «без изменений»."""
    dst = tmp_path / "content"
    shutil.copytree(CONTENT_DIR, dst)
    cast, bg = identity_rows(dst)
    rep = import_cast(write_xlsx(tmp_path / "fill.xlsx", cast, bg), content_root=dst, write=True)
    assert rep.ok, problems(rep)
    return dst


def chars(root: Path) -> dict:
    return {c.id: c for c in load_content(root).characters}


def identity_rows(root: Path) -> tuple[dict[str, list], dict[str, list]]:
    """Строки, повторяющие текущий контент: импорт без изменений. id → значения ячеек."""
    cast: dict[str, list] = {}
    bg: dict[str, list] = {}
    for ch in load_content(root).characters:
        if not ch.cast:
            continue
        if ch.importance == "background":
            if ch.father:
                name = ch.first
            else:
                name = f"{ch.first} {ch.last}"
                if ch.middle:
                    name = f"{ch.last} {ch.first} {ch.middle}"
            if "TODO" in name:
                name = "Алиса" if ch.father else "Пётр Сидоров" if ch.gender == "m" else "Анна Сидорова"
            bg[ch.id] = [ch.working, "", name, "Да"]
        else:
            name = f"{ch.last} {ch.first} {ch.middle}".strip()
            cast[ch.id] = [str(len(cast) + 1), ch.working, "", name, ch.age]
    return cast, bg


def write_xlsx(path: Path, cast: dict[str, list] | list[list], bg: dict[str, list] | list[list], *,
               cast_title: str = "Персонажи", bg_title: str = "Фоновые имена",
               cast_headers: list[str] = CAST_HEADERS, bg_headers: list[str] = BG_HEADERS,
               example_cast: list = EXAMPLE_CAST,
               extra_sheets: tuple[str, ...] = ()) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = cast_title
    ws.append(["Персонажи «За кулисами правды», версия 11.10.2026"])
    ws.append(cast_headers)
    ws.append(example_cast)
    for row in (cast.values() if isinstance(cast, dict) else cast):
        ws.append(row)
    if bg_title:
        ws = wb.create_sheet(bg_title)  # заголовки в 1-й строке, данные со 2-й, без примера
        ws.append(bg_headers)
        for row in (bg.values() if isinstance(bg, dict) else bg):
            ws.append(row)
    for title in extra_sheets:
        wb.create_sheet(title)
    wb.save(path)
    return path


def table(tmp_path: Path, content: Path, cast_patch: dict | None = None, bg_patch: dict | None = None,
          drop: tuple[str, ...] = (), **kw) -> Path:
    cast, bg = identity_rows(content)
    for cid, (name, *age) in (cast_patch or {}).items():
        cast[cid][3] = name
        if age:
            cast[cid][4] = age[0]
    for cid, name in (bg_patch or {}).items():
        bg[cid][2] = name
    for cid in drop:
        cast.pop(cid, None)
        bg.pop(cid, None)
    return write_xlsx(tmp_path / "cast.xlsx", cast, bg, **kw)


def problems(rep) -> str:
    return "\n".join(str(p) for p in rep.problems)


# ---------------------------------------------------------------- без изменений, сопоставление


def test_identity_table_no_changes(tmp_path, content):
    rep = import_cast(table(tmp_path, content), content_root=content)
    assert rep.ok, problems(rep)
    assert rep.changes == {}
    assert rep.new_text == rep.old_text
    assert {"kovalev", "sme"} <= set(rep.skipped)
    # строка-пример (3-я) не сопоставляется и не считается лишней
    assert not any("Пример" in str(w) for w in rep.warnings)


def test_matching_ignores_case_yo_and_spaces(tmp_path, content):
    cast, bg = identity_rows(content)
    cast["zavyalov"][1] = "  ЮРИЙ   завьялов "
    cast["stas"][1] = cast["stas"][1].replace("ё", "е")
    rep = import_cast(write_xlsx(tmp_path / "t.xlsx", cast, bg), content_root=content)
    assert rep.ok, problems(rep)


# ---------------------------------------------------------------- ключевой и фоновый


def test_key_character_import(tmp_path, content):
    path = table(tmp_path, content, cast_patch={"zavyalov": ("Петров Пётр Петрович", 70.0)})
    before = (content / "characters.yaml").read_bytes()
    rep = import_cast(path, content_root=content)
    assert rep.ok, problems(rep)
    assert rep.changes == {"zavyalov": {"last": ("Завьялов", "Петров"), "first": ("Юрий", "Пётр"),
                                        "middle": ("Викторович", "Петрович"), "age": (61, 70)}}
    assert '+    last: "Петров"' in rep.diff()
    assert (content / "characters.yaml").read_bytes() == before  # dry-run


def test_background_sheet_layout_header_row_1(tmp_path, content):
    # «Фоновые имена»: данные начинаются сразу со 2-й строки — первая строка данных не пропускается
    cast, bg = identity_rows(content)
    first = next(iter(bg))
    bg[first][2] = "Зоя Ким"
    rep = import_cast(write_xlsx(tmp_path / "t.xlsx", cast, bg), content_root=content)
    assert rep.ok, problems(rep)
    assert rep.changes[first]["first"][1] == "Зоя"


def test_background_character_import(tmp_path, content):
    path = table(tmp_path, content, bg_patch={"stagehand1": "Семён Лыков",
                                              "nina_actress": "Жукова Алла Борисовна"})
    rep = import_cast(path, content_root=content, write=True)
    assert rep.ok, problems(rep)
    c = chars(content)
    assert (c["stagehand1"].first, c["stagehand1"].last, c["stagehand1"].middle) == ("Семён", "Лыков", "")
    assert c["stagehand1"].gender == "m"  # «Имя Фамилия» — пол из content
    assert (c["nina_actress"].last, c["nina_actress"].first, c["nina_actress"].middle) == \
        ("Жукова", "Алла", "Борисовна")


# ---------------------------------------------------------------- father, familiar


def test_father_daughters_derived_from_new_father(tmp_path, content):
    path = table(tmp_path, content,
                 cast_patch={"kirill": ("Орлов Пётр Николаевич",), "marina": ("Ксения", 25)},
                 bg_patch={"marina_sister": "Алиса"})
    rep = import_cast(path, content_root=content, familiar={"voroncov": "Петя"}, write=True)
    assert rep.ok, problems(rep)
    c = chars(content)
    assert (c["marina"].last, c["marina"].first, c["marina"].middle, c["marina"].age) == \
        ("Орлова", "Ксения", "Петровна", 25)
    assert (c["marina_sister"].last, c["marina_sister"].first, c["marina_sister"].middle) == \
        ("Орлова", "Алиса", "Петровна")
    assert c["marina"].gender == "f"
    assert c["voroncov"].familiar == "Петя"


def test_father_child_full_fio_mismatch_is_problem(tmp_path, content):
    path = table(tmp_path, content, cast_patch={"marina": ("Сомова Марина Петровна",)})
    rep = import_cast(path, content_root=content)
    assert not rep.ok
    assert "отчество в таблице «Петровна»" in problems(rep)
    assert "«Кирилловна»" in problems(rep)


def test_familiar_problem_when_first_name_changes(tmp_path, content):
    path = table(tmp_path, content, cast_patch={"voroncov": ("Воронцов Пётр Сергеевич",)})
    before = (content / "characters.yaml").read_bytes()
    rep = import_cast(path, content_root=content, write=True)
    assert not rep.ok
    assert "уточните дружеское имя для voroncov (сейчас «Тёма»)" in problems(rep)
    assert (content / "characters.yaml").read_bytes() == before
    assert not rep.written


# ---------------------------------------------------------------- cast: false, пол


def test_cast_false_skipped_even_if_row_present(tmp_path, content):
    cast, bg = identity_rows(content)
    cast["kovalev"] = ["99", "Сержант Ковалев", "", "Кто-то", "abc"]  # мусор не проверяется
    rep = import_cast(write_xlsx(tmp_path / "t.xlsx", cast, bg), content_root=content)
    assert rep.ok, problems(rep)
    assert "kovalev" in rep.skipped and "kovalev" not in rep.changes
    assert not rep.warnings  # строка вымышленного персонажа не «лишняя»


def test_gender_from_patronymic(tmp_path, content):
    path = table(tmp_path, content, cast_patch={"gromov": ("Громов Михаил Анатольевич",)})
    rep = import_cast(path, content_root=content)
    assert rep.ok, problems(rep)
    assert rep.changes["gromov"]["gender"] == ("f", "m")
    assert any("пол меняется" in str(w) for w in rep.warnings)


def test_nonstandard_patronymic_is_problem(tmp_path, content):
    path = table(tmp_path, content, cast_patch={"zavyalov": ("Алиев Рустам Гасан-оглы",)})
    rep = import_cast(path, content_root=content)
    assert not rep.ok
    assert "нестандартное отчество «Гасан-оглы»" in problems(rep)
    assert "Персонажи!D" in problems(rep)  # ячейка с ФИО (колонка найдена по названию)


@pytest.mark.parametrize("middle,gender", [
    ("Викторович", "m"), ("Ильич", "m"), ("Никитич", "m"), ("Викторовна", "f"), ("Ильинична", "f"),
    ("Саввична", "f"), ("Кузьминична", "f"), ("Гасан-оглы", None), ("", None), (None, None),
    ("викторович", None), ("Петров", None),
])
def test_patronymic_gender(middle, gender):
    assert patronymic_gender(middle) == gender


# ---------------------------------------------------------------- ошибки таблицы


def test_missing_sheet(tmp_path, content):
    path = table(tmp_path, content, bg_title="", extra_sheets=("Лист2",))
    rep = import_cast(path, content_root=content)
    assert not rep.ok
    assert "нет листа «Фоновые имена» (есть: «Персонажи», «Лист2»)" in problems(rep)


def test_missing_column(tmp_path, content):
    headers = ["№", "Роль (рабочее имя)", "Комментарий", "ФИО", "Возраст в игре"]
    path = table(tmp_path, content, cast_headers=headers)
    rep = import_cast(path, content_root=content)
    assert not rep.ok
    assert "нет колонки «ФИО в документах»" in problems(rep)
    assert "Персонажи!2" in problems(rep)


def test_row_not_found(tmp_path, content):
    path = table(tmp_path, content, drop=("volkov",))
    rep = import_cast(path, content_root=content)
    assert not rep.ok
    assert "(volkov): в таблице нет строки «Артём Волков»" in problems(rep)


def test_two_part_fio_is_problem(tmp_path, content):
    path = table(tmp_path, content, cast_patch={"zavyalov": ("Петров Пётр",)})
    rep = import_cast(path, content_root=content)
    assert not rep.ok
    assert "«Петров Пётр» не разбирается: 2 части" in problems(rep)


def test_example_row_is_skipped(tmp_path, content):
    cast, bg = identity_rows(content)
    zav = cast.pop("zavyalov")
    zav[3] = "Петров Пётр Петрович"
    path = write_xlsx(tmp_path / "t.xlsx", cast, bg, example_cast=zav)
    rep = import_cast(path, content_root=content)
    assert not rep.ok
    text = problems(rep)
    assert "в таблице нет строки «Юрий Завьялов»" in text
    assert "строке 3 листа «Персонажи», но эта строка — пример" in text
    assert "zavyalov" not in rep.changes


def test_bad_age_and_duplicate_row(tmp_path, content):
    cast, bg = identity_rows(content)
    cast["volkov"][4] = "тридцать"
    rows = [*cast.values(), list(cast["stas"])]
    rep = import_cast(write_xlsx(tmp_path / "t.xlsx", rows, bg), content_root=content)
    text = problems(rep)
    assert "возраст «тридцать» — ожидается целое число лет" in text
    assert "встречается в таблице несколько раз" in text


def test_empty_background_name(tmp_path, content):
    # пусто + в content настоящее имя → остаётся (предупреждение); пусто + TODO(content) → проблема
    cast, bg = identity_rows(content)
    bg["rodionov_director"][2] = None
    rep = import_cast(write_xlsx(tmp_path / "t.xlsx", cast, bg), content_root=content)
    assert rep.ok, problems(rep)
    assert any("имя не заполнено — остаётся «Родионов Геннадий Павлович»" in str(w) for w in rep.warnings)

    raw = tmp_path / "raw"
    shutil.copytree(CONTENT_DIR, raw)
    chars = raw / "characters.yaml"  # имя врача СМП в content заполнено — вернуть заглушку для проверки
    chars.write_text(chars.read_text(encoding="utf-8").replace(
        'last: "Сафронов", first: "Артур", middle: "Вадимович"',
        'last: "TODO(content)", first: "TODO(content)", middle: ""'), encoding="utf-8")
    cast, bg = identity_rows(raw)
    bg["paramedic"][2] = None
    rep = import_cast(write_xlsx(tmp_path / "t2.xlsx", cast, bg), content_root=raw)
    assert "Фоновые имена!C" in problems(rep) and "не заполнено имя" in problems(rep)


def test_unknown_row_is_warning(tmp_path, content):
    cast, bg = identity_rows(content)
    cast["x"] = ["50", "Неизвестная роль", "", "Петров Пётр Петрович", 30]
    cast["total"] = ["", "Ролей заполнено:", "17 из 17", None, None]  # итоговая строка без имени — молча
    rep = import_cast(write_xlsx(tmp_path / "t.xlsx", cast, bg), content_root=content)
    assert rep.ok, problems(rep)
    assert [str(w) for w in rep.warnings] == [
        "t.xlsx: Персонажи!" + str(4 + len(cast) - 2)
        + ": строка «Неизвестная роль» не сопоставлена ни с одним персонажем (working в characters.yaml)"]


# ---------------------------------------------------------------- запись


def test_dry_run_cli_does_not_write(tmp_path, content, capsys):
    path = table(tmp_path, content, cast_patch={"zavyalov": ("Петров Пётр Петрович",)})
    before = (content / "characters.yaml").read_bytes()
    assert main(["import-cast", str(path), "--content", str(content)]) == 0
    out = capsys.readouterr().out
    assert "Dry-run: файл не изменён" in out
    assert "+    first: \"Пётр\"" in out
    assert (content / "characters.yaml").read_bytes() == before


def test_cli_problems_exit_1(tmp_path, content, capsys):
    path = table(tmp_path, content, cast_patch={"zavyalov": ("Петров Пётр",)})
    before = (content / "characters.yaml").read_bytes()
    assert main(["import-cast", str(path), "--content", str(content), "--write"]) == 1
    err = capsys.readouterr().err
    assert "Проблемы (1) — characters.yaml НЕ записан" in err
    assert (content / "characters.yaml").read_bytes() == before


def test_write_preserves_comments_and_passes_load_content(tmp_path, content, capsys):
    path = table(tmp_path, content, cast_patch={"zavyalov": ("Петров Пётр Петрович", 70),
                                                "belyaev": ("Авдеенко Михаил Ильич", 52)},
                 bg_patch={"rodionov_director": "Родионов Геннадий Павлович", "paramedic": "Олег Ким"})
    old = (content / "characters.yaml").read_text(encoding="utf-8")
    code = main(["import-cast", str(path), "--content", str(content), "--write"])
    assert code == 0, capsys.readouterr().err
    new = (content / "characters.yaml").read_text(encoding="utf-8")
    comments = [ln.split("#", 1)[1] for ln in old.splitlines() if "#" in ln]
    for c in comments:
        assert c in new
    # изменились только строки полей имени и возраста (и вставлена строка age у belyaev)
    changed = [ln for ln in difflib.ndiff(old.splitlines(), new.splitlines()) if ln[:2] in ("- ", "+ ")]
    field = re.compile(r"^[-+]\s{5}(last|first|middle|age|gender): ")
    assert changed and all(field.match(ln) or "{ id: paramedic," in ln for ln in changed), changed
    assert sum(ln.startswith("+ ") for ln in changed) == sum(ln.startswith("- ") for ln in changed) + 1
    c = load_content(content).characters  # проходит полную проверку
    by_id = {x.id: x for x in c}
    assert (by_id["belyaev"].last, by_id["belyaev"].age) == ("Авдеенко", 52)
    assert (by_id["paramedic"].first, by_id["paramedic"].last) == ("Олег", "Ким")
    assert [x.id for x in c] == [x.id for x in load_content(CONTENT_DIR).characters]


def test_write_rolls_back_when_load_content_fails(tmp_path, content, monkeypatch):
    import engine.importcast as ic

    path = table(tmp_path, content, cast_patch={"zavyalov": ("Петров Пётр Петрович",)})
    before = (content / "characters.yaml").read_bytes()
    calls = {"n": 0}
    real = ic.load_content

    def flaky(root):
        calls["n"] += 1
        if calls["n"] == 2:  # повторная загрузка после записи
            raise ContentError([])
        return real(root)

    monkeypatch.setattr(ic, "_validate", lambda *a: [])
    monkeypatch.setattr(ic, "load_content", flaky)
    rep = import_cast(path, content_root=content, write=True)
    assert not rep.ok and not rep.written
    assert "файл возвращён" in problems(rep)
    assert (content / "characters.yaml").read_bytes() == before


# ---------------------------------------------------------------- точечная правка YAML


def test_apply_edits_block_and_inline():
    text = (
        "# шапка\n"
        "characters:\n"
        "  # --- блок ---\n"
        "  - id: a\n"
        "    working: \"A, # не комментарий\"\n"
        "    last: \"Старый\"   # комментарий\n"
        "    first: 'Имя'\n"
        "    middle:\n"
        "    gender: m\n"
        "    forms:\n"
        "      last: { gen: \"Старого\" }\n"
        "  - { id: b, working: \"Г. Б, директор «X»\", last: \"TODO(content)\", first: \"TODO\", gender: f } # хвост\n"
        "other: 1\n"
    )
    new = apply_edits(text, {"a": {"last": "Новый", "middle": "Иванович", "age": 40},
                             "b": {"last": "Ким", "age": 7}})
    assert "# шапка" in new and "# --- блок ---" in new and "# хвост" in new
    assert '    last: "Новый"   # комментарий' in new
    assert '    middle: "Иванович"' in new
    assert "    gender: m\n    age: 40\n" in new
    assert '      last: { gen: "Старого" }' in new
    assert 'last: "Ким", first: "TODO", gender: f, age: 7 } # хвост' in new
    data = yaml.safe_load(new)
    assert data["characters"][0]["last"] == "Новый" and data["characters"][0]["first"] == "Имя"
    assert data["characters"][1]["working"] == "Г. Б, директор «X»"
    assert data["other"] == 1


def test_report_format_lists_problems_and_diff(tmp_path, content):
    path = table(tmp_path, content, cast_patch={"zavyalov": ("Петров Пётр Петрович",),
                                                "volkov": ("Волков Артём",)})
    rep = import_cast(path, content_root=content)
    out, err = format_report(rep, write=False)
    assert "Предварительный diff" in out and '+    last: "Петров"' in out
    assert "Проблемы (1)" in err and "Волков Артём" in err

