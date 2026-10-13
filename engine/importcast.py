"""``python -m engine import-cast <xlsx>`` — настоящие имена из таблицы персонажей (P09 п.1).

Таблица версии от 11.10.2026 (``SHEETS``): лист «Персонажи» — заголовки во 2-й строке,
3-я — пример (пропускается), данные с 4-й; лист «Фоновые имена» — так же. Колонки ищутся
по названию заголовка (``engine.normalize.base``: без учёта регистра, ё/е, пунктуации и
лишних пробелов), строка ↔ персонаж — по ``working`` (то же сравнение).

Результат — новая версия ``characters.yaml``: меняются только значения ``last``, ``first``,
``middle``, ``age``, ``gender`` (и ``familiar`` — только по явному ``--familiar ID=Имя``)
точечной построчной правкой; комментарии, порядок и прочие поля сохраняются. Любая
проблема — отчёт и отказ записывать. Новая версия проверяется ``load_content`` на копии
``content/`` до записи и ещё раз после записи (ошибка → откат).
"""

from __future__ import annotations

import difflib
import json
import re
import shutil
import tempfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import yaml

from engine.content import CONTENT_DIR, TODO_MARK, ContentError, load_content
from engine.models import Character
from engine.names import DeclensionError, patronymic_for, patronymic_gender, surname_for
from engine.normalize import base
from engine.util import format_loc

CHARACTERS_FILE = "characters.yaml"
EDITABLE = ("last", "first", "middle", "age", "gender", "familiar")
# Куда вставлять поле, если его нет в записи: после первого найденного из списка.
INSERT_AFTER: dict[str, tuple[str, ...]] = {
    "last": ("importance", "role", "working", "id"),
    "first": ("last", "importance", "role", "working", "id"),
    "middle": ("first", "last", "id"),
    "gender": ("nick", "middle", "first", "last", "id"),
    "age": ("gender", "nick", "middle", "first", "last", "id"),
    "familiar": ("age", "gender", "middle", "first", "last", "id"),
}


# ---------------------------------------------------------------- формат таблицы


@dataclass(frozen=True)
class SheetSpec:
    kind: str  # cast — «Персонажи» (ФИО + возраст); background — «Фоновые имена»
    title: str
    header_row: int
    example_row: int
    first_row: int
    key: str  # колонка с рабочим именем (= working)
    name: str  # колонка с именем
    age: str | None = None  # колонка с возрастом


# Таблица версии от 11.10.2026.
CAST_SHEET = SheetSpec("cast", "Персонажи", header_row=2, example_row=3, first_row=4,
                       key="Роль (рабочее имя)", name="ФИО в документах", age="Возраст в игре")
BACKGROUND_SHEET = SheetSpec("background", "Фоновые имена", header_row=2, example_row=3, first_row=4,
                             key="Кто", name="Имя (заполни)")
SHEETS: tuple[SheetSpec, ...] = (CAST_SHEET, BACKGROUND_SHEET)
BACKGROUND_IMPORTANCE = "background"  # их строка ищется сначала на листе «Фоновые имена»


@dataclass(frozen=True)
class Issue:
    where: str
    msg: str

    def __str__(self) -> str:
        return f"{self.where}: {self.msg}" if self.where else self.msg


@dataclass
class TableRow:
    sheet: SheetSpec
    row: int
    working: str
    name: Any
    name_cell: str
    age: Any = None
    age_cell: str = ""

    def where(self, xlsx: str, cell: str | None = None) -> str:
        col = {self.name_cell: self.sheet.name, self.age_cell: self.sheet.age}.get(cell or self.name_cell)
        ref = f"{self.sheet.title}!{cell or self.name_cell}"
        return f"{xlsx}: {ref} («{col}», «{self.working}»)"


@dataclass
class ImportReport:
    path: Path
    xlsx: Path
    problems: list[Issue] = field(default_factory=list)
    warnings: list[Issue] = field(default_factory=list)
    skipped: list[str] = field(default_factory=list)  # cast: false
    changes: dict[str, dict[str, tuple[Any, Any]]] = field(default_factory=dict)
    old_text: str = ""
    new_text: str = ""
    written: bool = False

    @property
    def ok(self) -> bool:
        return not self.problems

    def diff(self) -> str:
        label = f"{self.path.parent.name}/{self.path.name}"
        return "".join(difflib.unified_diff(
            self.old_text.splitlines(keepends=True), self.new_text.splitlines(keepends=True),
            fromfile=f"a/{label}", tofile=f"b/{label}",
        ))


class _Fatal(Exception):
    """Дальше разбирать нечего (нет файла, листа, колонки)."""


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    return " ".join(str(value).split())


def _read_sheets(xlsx: Path, rep: ImportReport) -> dict[str, list[TableRow]]:
    import openpyxl
    from openpyxl.utils import get_column_letter

    name = xlsx.name
    if not xlsx.is_file():
        raise _Fatal(f"{xlsx}: файл не найден")
    try:
        wb = openpyxl.load_workbook(xlsx, data_only=True)
    except Exception as e:  # noqa: BLE001 — openpyxl бросает что угодно на битом файле
        raise _Fatal(f"{xlsx}: не удалось прочитать таблицу .xlsx ({type(e).__name__}: {e})") from None

    titles = {base(t): t for t in wb.sheetnames}
    out: dict[str, list[TableRow]] = {}
    fatal = False
    for spec in SHEETS:
        title = titles.get(base(spec.title))
        if title is None:
            have = ", ".join(f"«{t}»" for t in wb.sheetnames)
            rep.problems.append(Issue(name, f"нет листа «{spec.title}» (есть: {have})"))
            fatal = True
            continue
        ws = wb[title]
        headers: dict[str, list[int]] = {}
        for col, cell in enumerate(next(ws.iter_rows(min_row=spec.header_row, max_row=spec.header_row,
                                                     values_only=True), ()), start=1):
            key = base(_cell_text(cell))
            if key:
                headers.setdefault(key, []).append(col)
        cols: dict[str, int] = {}
        for what in (spec.key, spec.name, spec.age):
            if what is None:
                continue
            found = headers.get(base(what), [])
            if len(found) != 1:
                if found:
                    letters = ", ".join(get_column_letter(c) for c in found)
                    msg = f"колонка «{what}» встречается несколько раз ({letters})"
                else:
                    have = ", ".join(f"«{_cell_text(ws.cell(spec.header_row, c).value)}»"
                                     for cs in headers.values() for c in cs) or "пусто"
                    msg = f"нет колонки «{what}» в заголовках (строка {spec.header_row}: {have})"
                rep.problems.append(Issue(f"{name}: {title}!{spec.header_row}", msg))
                fatal = True
                continue
            cols[what] = found[0]
        if len(cols) != len([c for c in (spec.key, spec.name, spec.age) if c]):
            continue

        rows: list[TableRow] = []
        kc, nc = cols[spec.key], cols[spec.name]
        ac = cols.get(spec.age) if spec.age else None
        for r in range(spec.example_row, ws.max_row + 1):
            working = _cell_text(ws.cell(r, kc).value)
            value = ws.cell(r, nc).value
            age = ws.cell(r, ac).value if ac else None
            row = TableRow(spec, r, working, value, f"{get_column_letter(nc)}{r}", age,
                           f"{get_column_letter(ac)}{r}" if ac else "")
            if r < spec.first_row:
                rows.append(row)  # пример: сопоставляется только для предупреждения
                continue
            if not working:
                if _cell_text(value) or _cell_text(age):
                    rep.warnings.append(Issue(f"{name}: {title}!{get_column_letter(kc)}{r}",
                                              f"пустая колонка «{spec.key}» — строка пропущена"))
                continue
            rows.append(row)
        out[spec.kind] = rows
    if fatal:
        raise _Fatal("")
    return out


# ---------------------------------------------------------------- разбор имени и возраста


@dataclass
class Parsed:
    first: str
    last: str | None = None
    middle: str | None = None
    gender: str | None = None  # по отчеству


def _ch_where(i: int, ch: Character, part: str | None = None) -> str:
    loc = ("characters", i, part) if part else ("characters", i)
    return f"{CHARACTERS_FILE}: {format_loc(loc)} ({ch.id})"


_PATRONYMIC_HINT = ("пол определяется по отчеству: -вич → m; -вна/-чна/-ична → f "
                    "(одно слово с заглавной буквы)")


def _parse_name(row: TableRow, ch: Character, xlsx: str, rep: ImportReport) -> Parsed | None:
    text = _cell_text(row.name)
    where = row.where(xlsx)
    if not text:
        rep.problems.append(Issue(where, "не заполнено имя"))
        return None
    parts = text.split(" ")
    child = ch.father is not None
    n = len(parts)
    if n == 3:
        last, first, middle = parts
        gender = patronymic_gender(middle)
        if gender is None:
            rep.problems.append(Issue(where, f"нестандартное отчество «{middle}» в «{text}» — {_PATRONYMIC_HINT}; "
                                             "уточните пол и отчество"))
            return None
        return Parsed(first, last, middle, gender)
    if child and n == 1:
        return Parsed(parts[0])
    if row.sheet.kind == "background" and n == 2:
        return Parsed(parts[0], parts[1])  # «Имя Фамилия»; пол — из content
    if child:
        expect = ("только имя, «Имя Фамилия» или «Фамилия Имя Отчество»" if row.sheet.kind == "background"
                  else "только имя или «Фамилия Имя Отчество»")
    elif row.sheet.kind == "background":
        expect = "«Имя Фамилия» или «Фамилия Имя Отчество»"
    else:
        expect = "«Фамилия Имя Отчество» (3 части)"
    rep.problems.append(Issue(where, f"«{text}» не разбирается: {n} част{'ь' if n == 1 else 'и'}, ожидается {expect}"))
    return None


def _parse_age(row: TableRow, xlsx: str, rep: ImportReport) -> int | None:
    v = row.age
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, bool):
        n = None
    elif isinstance(v, int):
        n = v
    elif isinstance(v, float) and v.is_integer():
        n = int(v)
    elif isinstance(v, str) and v.strip().isdigit():
        n = int(v.strip())
    else:
        n = None
    if n is None or n < 0:
        rep.problems.append(Issue(row.where(xlsx, row.age_cell),
                                  f"возраст «{_cell_text(v)}» — ожидается целое число лет"))
        return None
    return n


# ---------------------------------------------------------------- построение новых значений


def _source(ch: Character, new: dict[str, Any], part: str) -> str:
    """Часть имени отца, от которой выводят: ручная forms.<part>.nom важнее значения."""
    nom = ((ch.forms or {}).get(part) or {}).get("nom")
    return nom if nom is not None else new[part]


def _build(characters: list[Character], table: dict[str, list[TableRow]], xlsx: str,
           rep: ImportReport, familiar: dict[str, str]) -> dict[str, dict[str, Any]]:
    """id → новые значения EDITABLE (только у сопоставленных и без проблем)."""
    by_key: dict[str, dict[str, list[TableRow]]] = {}
    examples: dict[str, TableRow] = {}
    for kind, rows in table.items():
        for row in rows:
            k = base(row.working)
            if row.row < row.sheet.first_row:
                if k:
                    examples[k] = row
                continue
            by_key.setdefault(k, {}).setdefault(kind, []).append(row)

    index = {ch.id: i for i, ch in enumerate(characters)}
    by_id = {ch.id: ch for ch in characters}
    keys: dict[str, list[str]] = {}
    for ch in characters:
        keys.setdefault(base(ch.working), []).append(ch.id)
    for k, ids in keys.items():
        if len(ids) > 1:
            rep.problems.append(Issue(CHARACTERS_FILE, f"одинаковое рабочее имя (working) у {', '.join(ids)} — "
                                                       "строку таблицы не сопоставить"))

    used: set[int] = set()
    matched: dict[str, TableRow] = {}
    for i, ch in enumerate(characters):
        if not ch.cast:
            rep.skipped.append(ch.id)
            continue
        k = base(ch.working)
        if len(keys.get(k, [])) > 1:
            continue
        found = by_key.get(k, {})
        order = (("background", "cast") if ch.importance == BACKGROUND_IMPORTANCE else ("cast", "background"))
        rows = [r for kind in order for r in found.get(kind, [])]
        for r in rows:
            used.add(id(r))
        if not rows:
            sheet = BACKGROUND_SHEET if ch.importance == BACKGROUND_IMPORTANCE else CAST_SHEET
            msg = f"в таблице нет строки «{ch.working}» (лист «{sheet.title}», колонка «{sheet.key}»)"
            ex = examples.get(k)
            if ex is not None:
                msg += (f"; она есть в строке {ex.row} листа «{ex.sheet.title}», но эта строка — пример "
                        f"и пропускается (данные с {ex.sheet.first_row}-й)")
            rep.problems.append(Issue(_ch_where(i, ch, "working"), msg))
            continue
        if len(rows) > 1:
            refs = ", ".join(f"{r.sheet.title}!{r.row}" for r in rows)
            rep.problems.append(Issue(_ch_where(i, ch, "working"),
                                      f"«{ch.working}» встречается в таблице несколько раз ({refs})"))
            continue
        matched[ch.id] = rows[0]

    for kind, rows in table.items():
        for row in rows:
            if row.row >= row.sheet.first_row and id(row) not in used:
                cast_false = [c for c in keys.get(base(row.working), []) if not by_id[c].cast]
                if cast_false:
                    continue  # строка вымышленного персонажа — не ищется и не считается лишней
                rep.warnings.append(Issue(f"{xlsx}: {row.sheet.title}!{row.row}",
                                          f"строка «{row.working}» не сопоставлена ни с одним персонажем "
                                          f"(working в {CHARACTERS_FILE})"))

    parsed: dict[str, tuple[Parsed, int | None]] = {}
    for cid, row in matched.items():
        ch = by_id[cid]
        p = _parse_name(row, ch, xlsx, rep)
        age = _parse_age(row, xlsx, rep) if row.sheet.age else None
        if p is not None:
            parsed[cid] = (p, age)

    new: dict[str, dict[str, Any]] = {}

    def current(ch: Character) -> dict[str, Any]:
        return {"last": ch.last, "first": ch.first, "middle": ch.middle or "", "gender": ch.gender,
                "age": ch.age, "familiar": ch.familiar}

    # 1) без father
    for cid, (p, age) in parsed.items():
        ch = by_id[cid]
        if ch.father is not None:
            continue
        vals = current(ch)
        vals.update(last=p.last, first=p.first, middle=p.middle or "", gender=p.gender or ch.gender)
        if age is not None:
            vals["age"] = age
        new[cid] = vals

    # 2) с father — от отца из новой версии (рекурсивно, отец раньше ребёнка)
    resolving: set[str] = set()

    def father_values(fid: str) -> dict[str, Any] | None:
        f = by_id.get(fid)
        if f is None:
            return None
        if fid in new:
            return new[fid]
        if f.father is not None and fid in parsed:
            return resolve(fid)
        if f.cast and f.id not in parsed:
            return None  # у отца своя проблема — уже в списке
        return current(f)

    def resolve(cid: str) -> dict[str, Any] | None:
        if cid in new:
            return new[cid]
        if cid in resolving:
            rep.problems.append(Issue(_ch_where(index[cid], by_id[cid], "father"), "цикл father"))
            return None
        resolving.add(cid)
        try:
            return _resolve_child(cid)
        finally:
            resolving.discard(cid)

    def _resolve_child(cid: str) -> dict[str, Any] | None:
        ch = by_id[cid]
        i = index[cid]
        p, age = parsed[cid]
        father = by_id.get(ch.father or "")
        fv = father_values(ch.father or "")
        if father is None or fv is None:
            return None
        if fv["gender"] != "m":
            rep.problems.append(Issue(_ch_where(i, ch, "father"),
                                      f"отец «{father.id}» по новым данным женского пола — "
                                      "фамилию и отчество не вывести"))
            return None
        gender = p.gender or ch.gender
        ok = True
        out = current(ch)
        out.update(first=p.first, gender=gender)
        if age is not None:
            out["age"] = age
        old_fv = current(father)
        derive = {
            "last": ("last", lambda v: surname_for(v, gender), "фамилия", "фамилии"),
            "middle": ("first", lambda v: patronymic_for(v, gender), "отчество", "имени"),
        }
        for part, (src_part, rule, what, src_what) in derive.items():
            src = _source(father, fv, src_part)
            manual = ((ch.forms or {}).get(part) or {}).get("nom")
            if manual is not None:
                value = manual
                if src != _source(father, old_fv, src_part):
                    rep.warnings.append(Issue(_ch_where(i, ch, f"forms.{part}.nom"),
                                              f"ручная форма «{manual}» сохранена; проверьте её для нового "
                                              f"{src_what} отца «{src}»"))
            elif TODO_MARK in src:
                rep.problems.append(Issue(_ch_where(i, ch, part),
                                          f"{what} не вывести: у отца «{father.id}» {src_part} — {TODO_MARK}"))
                ok = False
                continue
            else:
                try:
                    value = rule(src)
                except DeclensionError as e:
                    rep.problems.append(Issue(_ch_where(i, ch, part),
                                              f"{e}. Нужно отчество «{ch.id}» от имени отца «{src}»: "
                                              f"впишите его в {CHARACTERS_FILE} → {ch.id}: middle и forms.middle.nom"))
                    ok = False
                    continue
            given = p.last if part == "last" else p.middle
            if given is not None and given != value:
                rep.problems.append(Issue(matched[cid].where(xlsx),
                                          f"{what} в таблице «{given}», а от отца «{father.id}» выводится "
                                          f"«{value}» — уточните (или задайте forms.{part}.nom)"))
                ok = False
            out[part] = value
        if not ok:
            return None
        new[cid] = out
        return out

    for cid in parsed:
        if by_id[cid].father is not None:
            resolve(cid)

    # familiar, forms: только проверки
    for cid, vals in new.items():
        ch = by_id[cid]
        i = index[cid]
        if cid in familiar:
            vals["familiar"] = familiar[cid]
        elif ch.familiar is not None and vals["first"] != ch.first:
            rep.problems.append(Issue(_ch_where(i, ch, "familiar"),
                                      f"уточните дружеское имя для {cid} (сейчас «{ch.familiar}»): имя меняется "
                                      f"«{ch.first}» → «{vals['first']}»; укажите --familiar {cid}=<имя>"))
        for part in ("last", "first", "middle"):
            if ch.father is not None and part in ("last", "middle"):
                continue
            manual = (ch.forms or {}).get(part)
            if manual and vals[part] != getattr(ch, part):
                shown = ", ".join(f"{k}: «{v}»" for k, v in manual.items())
                rep.warnings.append(Issue(_ch_where(i, ch, f"forms.{part}"),
                                          f"ручные формы ({shown}) заданы для «{getattr(ch, part)}», а новое "
                                          f"значение «{vals[part]}» — проверьте их (python -m engine names)"))
        if ch.middle and not vals["middle"]:
            rep.warnings.append(Issue(_ch_where(i, ch, "middle"),
                                      f"отчество «{ch.middle}» удаляется: в таблице имя без отчества"))
        if vals["gender"] != ch.gender:
            rep.warnings.append(Issue(_ch_where(i, ch, "gender"),
                                      f"пол меняется: {ch.gender} → {vals['gender']} (по отчеству «{vals['middle']}»)"))
    for cid in familiar:
        if cid not in by_id:
            rep.problems.append(Issue("--familiar", f"нет персонажа «{cid}» в {CHARACTERS_FILE}"))
        elif cid not in new and cid not in matched:
            rep.problems.append(Issue("--familiar", f"персонаж «{cid}» не импортируется из таблицы"))
    return new


# ---------------------------------------------------------------- точечная правка YAML


class YamlEditError(Exception):
    pass


def _split_scalar(s: str, flow: bool = False) -> tuple[str, str, str]:
    """``s`` — текст после ``key:`` → (пробелы, скаляр, хвост с комментарием).

    ``flow`` — внутри ``{ … }``: простой скаляр кончается и на «,», «}», «]»."""
    i = len(s) - len(s.lstrip(" \t"))
    lead = s[:i]
    if i < len(s) and s[i] == '"':
        j = i + 1
        while j < len(s) and s[j] != '"':
            j += 2 if s[j] == "\\" else 1
        if j >= len(s):
            raise YamlEditError("незакрытая кавычка")
        end = j + 1
    elif i < len(s) and s[i] == "'":
        j = i + 1
        while j < len(s):
            if s[j] == "'":
                if j + 1 < len(s) and s[j + 1] == "'":
                    j += 2
                    continue
                break
            j += 1
        if j >= len(s):
            raise YamlEditError("незакрытая кавычка")
        end = j + 1
    else:
        m = re.search(r"\s#|[,}\]]" if flow else r"\s#", s[i:])
        end = i + (m.start() if m else len(s) - i)
        while end > i and s[end - 1] in " \t":
            end -= 1
    return lead, s[i:end], s[end:]


def _fmt(field_name: str, value: Any) -> str:
    if field_name == "age":
        return "null" if value is None else str(int(value))
    if field_name == "gender":
        return str(value)
    return json.dumps(value, ensure_ascii=False)  # двойные кавычки JSON — валидный YAML


_KEY = re.compile(r"([A-Za-z_][A-Za-z0-9_]*)\s*:(.*)$", re.S)


def _flow_items(body: str) -> list[str]:
    """Содержимое ``{ … }`` → элементы по запятым верхнего уровня (кавычки и скобки учтены)."""
    items, depth, cur, i = [], 0, [], 0
    while i < len(body):
        c = body[i]
        if c in "\"'":
            _, tok, _ = _split_scalar(body[i:])
            cur.append(tok)
            i += len(tok)
            continue
        if c in "[{":
            depth += 1
        elif c in "]}":
            depth -= 1
        if c == "," and depth == 0:
            items.append("".join(cur))
            cur = []
        else:
            cur.append(c)
        i += 1
    items.append("".join(cur))
    return items


def _flow_end(line: str, start: int) -> int:
    """Индекс закрывающей ``}`` для ``{`` в позиции ``start``."""
    depth, i = 0, start
    while i < len(line):
        c = line[i]
        if c in "\"'":
            _, tok, _ = _split_scalar(line[i:])
            i += len(tok)
            continue
        if c in "[{":
            depth += 1
        elif c in "]}":
            depth -= 1
            if depth == 0:
                return i
        i += 1
    raise YamlEditError("запись { … } должна умещаться в одну строку")


def _edit_flow(line: str, brace: int, cid: str, edits: dict[str, Any]) -> str:
    end = _flow_end(line, brace)
    items = _flow_items(line[brace + 1:end])
    keys: list[str | None] = []
    for it in items:
        m = _KEY.match(it.strip())
        keys.append(m.group(1) if m else None)
    for name, value in edits.items():
        if name in keys:
            k = keys.index(name)
            it = items[k]
            colon = it.index(":")
            lead, _tok, rest = _split_scalar(it[colon + 1:], flow=True)
            items[k] = it[: colon + 1] + (lead or " ") + _fmt(name, value) + rest
        else:
            anchor = next((a for a in INSERT_AFTER[name] if a in keys), None)
            if anchor is None:
                raise YamlEditError(f"не к чему привязать новое поле {name} у «{cid}»")
            k = keys.index(anchor) + 1
            prev = items[k - 1]
            tail = prev[len(prev.rstrip()):]  # пробел перед «}» остаётся в конце
            items[k - 1] = prev.rstrip()
            items.insert(k, f" {name}: {_fmt(name, value)}{tail}")
            keys.insert(k, name)
    return line[: brace + 1] + ",".join(items) + line[end:]


def _items(lines: list[str]) -> dict[str, tuple[int, int, int]]:
    """id → (первая строка, конец, отступ «-») для записей списка characters."""
    top = next((n for n, ln in enumerate(lines) if re.match(r"^characters\s*:", ln)), None)
    if top is None:
        raise YamlEditError("нет ключа characters на верхнем уровне")
    starts: list[tuple[int, int]] = []
    dash: int | None = None
    end_all = len(lines)
    for n in range(top + 1, len(lines)):
        ln = lines[n]
        if not ln.strip() or ln.lstrip().startswith("#"):
            continue
        ind = len(ln) - len(ln.lstrip(" "))
        if ind == 0 and not ln.startswith("-"):
            end_all = n
            break
        if re.match(r"^ *- ", ln) and (dash is None or ind == dash):
            dash = ind
            starts.append((n, ind))
    out: dict[str, tuple[int, int, int]] = {}
    for k, (s, d) in enumerate(starts):
        e = starts[k + 1][0] if k + 1 < len(starts) else end_all
        first = lines[s][d + 2:]
        if first.lstrip().startswith("{"):
            body = first[first.index("{"):]
            m = re.search(r"(?:^|[{,])\s*id\s*:", body)
            if not m:
                continue
            _, tok, _ = _split_scalar(body[m.end():], flow=True)
        else:
            tok = None
            for n in range(s, e):
                pre = lines[n][: d + 2]
                if pre not in (" " * d + "- ", " " * (d + 2)):
                    continue
                m = _KEY.match(lines[n][d + 2:])
                if m and m.group(1) == "id":
                    _, tok, _ = _split_scalar(m.group(2))
                    break
            if tok is None:
                continue
        cid = yaml.safe_load(tok)
        if cid in out:
            raise YamlEditError(f"id «{cid}» встречается дважды")
        out[str(cid)] = (s, e, d)
    return out


def apply_edits(text: str, edits: dict[str, dict[str, Any]]) -> str:
    """Точечно заменить значения полей у записей по id; комментарии и порядок сохраняются."""
    lines = text.split("\n")
    items = _items(lines)
    # снизу вверх: вставки строк не сдвигают ещё не обработанные записи
    for cid in sorted(edits, key=lambda c: -items[c][0] if c in items else 0):
        if cid not in items:
            raise YamlEditError(f"запись id «{cid}» не найдена")
        s, e, d = items[cid]
        first = lines[s][d + 2:]
        if first.lstrip().startswith("{"):
            lines[s] = _edit_flow(lines[s], lines[s].index("{", d), cid, edits[cid])
            continue
        fields: dict[str, int] = {}
        for n in range(s, e):
            pre = lines[n][: d + 2]
            if pre not in (" " * d + "- ", " " * (d + 2)):
                continue
            m = _KEY.match(lines[n][d + 2:])
            if m:
                fields.setdefault(m.group(1), n)
        inserts: list[tuple[int, str]] = []
        for name, value in edits[cid].items():
            if name in fields:
                n = fields[name]
                head, after = lines[n].split(":", 1)
                lead, _tok, rest = _split_scalar(after)
                lines[n] = f"{head}:{lead or ' '}{_fmt(name, value)}{rest}"
            else:
                anchor = next((a for a in INSERT_AFTER[name] if a in fields), None)
                if anchor is None:
                    raise YamlEditError(f"не к чему привязать новое поле {name} у «{cid}»")
                inserts.append((fields[anchor], f"{' ' * (d + 2)}{name}: {_fmt(name, value)}"))
        for n, new_line in sorted(inserts, key=lambda x: -x[0]):
            lines.insert(n + 1, new_line)
    return "\n".join(lines)


# ---------------------------------------------------------------- команда


def _validate(root: Path, text: str, newline: str) -> list[Issue]:
    """Проверить новую версию load_content на копии content/."""
    with tempfile.TemporaryDirectory() as tmp:
        copy = Path(tmp) / root.name
        shutil.copytree(root, copy)
        _write(copy / CHARACTERS_FILE, text, newline)
        try:
            load_content(copy)
        except ContentError as e:
            return [Issue("", f"после импорта контент не проходит проверку: {err}") for err in e.errors]
    return []


def _read(path: Path) -> tuple[str, str]:
    raw = path.read_bytes().decode("utf-8")
    newline = "\r\n" if "\r\n" in raw else "\n"
    return raw.replace("\r\n", "\n"), newline


def _write(path: Path, text: str, newline: str) -> None:
    with path.open("w", encoding="utf-8", newline=newline) as fh:
        fh.write(text)


def import_cast(xlsx: Path | str, *, content_root: Path | str = CONTENT_DIR, write: bool = False,
                familiar: dict[str, str] | None = None) -> ImportReport:
    root = Path(content_root)
    xlsx = Path(xlsx)
    rep = ImportReport(path=root / CHARACTERS_FILE, xlsx=xlsx)
    try:
        content = load_content(root)
    except ContentError as e:
        rep.problems.append(Issue("", "текущий контент не проходит проверку — сначала исправьте его"))
        rep.problems.extend(Issue("", str(err)) for err in e.errors)
        return rep
    rep.old_text, newline = _read(rep.path)
    rep.new_text = rep.old_text
    try:
        table = _read_sheets(xlsx, rep)
    except _Fatal as e:
        if str(e):
            rep.problems.append(Issue("", str(e)))
        return rep

    characters = content.characters
    new = _build(characters, table, xlsx.name, rep, dict(familiar or {}))
    edits: dict[str, dict[str, Any]] = {}
    for ch in characters:
        vals = new.get(ch.id)
        if vals is None:
            continue
        old = {"last": ch.last, "first": ch.first, "middle": ch.middle or "", "gender": ch.gender,
               "age": ch.age, "familiar": ch.familiar}
        diff = {k: vals[k] for k in EDITABLE if vals[k] != old[k] and not (k == "age" and vals[k] is None)}
        if diff:
            edits[ch.id] = diff
            rep.changes[ch.id] = {k: (old[k], v) for k, v in diff.items()}

    try:
        rep.new_text = apply_edits(rep.old_text, edits)
    except YamlEditError as e:
        rep.problems.append(Issue(CHARACTERS_FILE, f"не удалось точечно изменить файл: {e}"))
        rep.new_text = rep.old_text
        return rep

    # самопроверка правки: данные = старые + ровно наши изменения
    expected = yaml.safe_load(rep.old_text)
    for item in expected["characters"]:
        item.update(edits.get(item["id"], {}))
    if yaml.safe_load(rep.new_text) != expected:
        rep.problems.append(Issue(CHARACTERS_FILE, "внутренняя ошибка: точечная правка изменила не те данные"))
        return rep

    if rep.problems:
        return rep
    if edits:
        rep.problems.extend(_validate(root, rep.new_text, newline))
    if rep.problems or not write or not edits:
        return rep

    _write(rep.path, rep.new_text, newline)
    try:
        load_content(root)
    except ContentError as e:
        _write(rep.path, rep.old_text, newline)
        rep.problems.append(Issue(CHARACTERS_FILE, "после записи контент не прошёл проверку — файл возвращён"))
        rep.problems.extend(Issue("", str(err)) for err in e.errors)
        return rep
    rep.written = True
    return rep


def format_report(rep: ImportReport, write: bool) -> tuple[str, str]:
    """(stdout, stderr) отчёта команды."""
    out: list[str] = [f"Таблица: {rep.xlsx}", f"Файл: {rep.path}"]
    if rep.skipped:
        out.append(f"Пропущены (cast: false): {', '.join(rep.skipped)}")
    if rep.warnings:
        out.append(f"Предупреждения ({len(rep.warnings)}):")
        out.extend(f"  {w}" for w in rep.warnings)
    if rep.changes:
        out.append(f"Изменения ({len(rep.changes)} персонаж.):")
        for cid, ch in rep.changes.items():
            parts = "; ".join(f"{k}: {_show(a)} → {_show(b)}" for k, (a, b) in ch.items())
            out.append(f"  {cid}: {parts}")
        d = rep.diff()
        if d:
            out.append("")
            if rep.problems:
                out.append("Предварительный diff (без персонажей с проблемами):")
            out.append(d.rstrip("\n"))
    else:
        out.append("Изменений нет: characters.yaml уже совпадает с таблицей.")
    err: list[str] = []
    if rep.problems:
        err.append(f"Проблемы ({len(rep.problems)}) — {CHARACTERS_FILE} НЕ записан:")
        err.extend(f"  {p}" for p in rep.problems)
    elif rep.written:
        out.append(f"Записано: {rep.path} (проверка load_content пройдена). "
                   "Дальше: python -m engine names — вычитать падежи.")
    elif rep.changes:
        out.append("Dry-run: файл не изменён. Записать: python -m engine import-cast <xlsx> --write")
    return "\n".join(out), "\n".join(err)


def _show(v: Any) -> str:
    return "—" if v is None or v == "" else f"«{v}»" if isinstance(v, str) else str(v)
